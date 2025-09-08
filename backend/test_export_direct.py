#!/usr/bin/env python3

"""
直接测试导出功能
"""

import sys
import os
import traceback
from io import BytesIO

# 添加当前目录到路径
sys.path.append(os.path.dirname(os.path.abspath(__file__)))

from database import get_db
import crud
import models
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime

def test_export_direct():
    """直接测试导出功能"""
    print("开始测试导出功能...")
    
    db = next(get_db())
    try:
        # 获取汇总数据
        print("1. 获取汇总数据...")
        summary_data = crud.get_room_summary(db, None)
        rooms = summary_data["rooms"]
        print(f"获取到 {len(rooms)} 个房间")
        
        if not rooms:
            print("没有房间数据，无法测试导出")
            return
        
        # 查看第一个房间的字段
        first_room = rooms[0]
        print(f"第一个房间的字段: {list(first_room.keys())}")
        
        # 应用筛选条件（无筛选）
        print("2. 应用筛选条件...")
        filtered_rooms = []
        for room in rooms:
            filtered_rooms.append(room)
        
        print(f"筛选后有 {len(filtered_rooms)} 个房间")
        
        # 创建Excel工作簿
        print("3. 创建Excel工作簿...")
        wb = Workbook()
        ws = wb.active
        ws.title = "房间信息汇总"
        
        # 设置标题行
        headers = [
            "楼栋", "房间号", "户主姓名", "手机号码", "整改状态", 
            "交付状态", "签约状态", "待验收", "预计交付时间", "信件状态"
        ]
        
        # 写入标题行
        print("4. 写入标题行...")
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        
        # 写入数据行
        print("5. 写入数据行...")
        for row_idx, room in enumerate(filtered_rooms, 2):
            print(f"处理第 {row_idx-1} 行数据: {room.get('building_unit')}-{room.get('room_number')}")
            
            ws.cell(row=row_idx, column=1, value=room.get("building_unit", ""))
            ws.cell(row=row_idx, column=2, value=room.get("room_number", ""))
            ws.cell(row=row_idx, column=3, value=room.get("owner_name", "") or "未录入")
            ws.cell(row=row_idx, column=4, value=room.get("owner_phone", "") or "未录入")
            ws.cell(row=row_idx, column=5, value=room.get("status", ""))
            ws.cell(row=row_idx, column=6, value=room.get("delivery_status", ""))
            ws.cell(row=row_idx, column=7, value=room.get("contract_status", ""))
            ws.cell(row=row_idx, column=8, value=room.get("pending_issues_count", 0))
            
            # 处理日期格式
            expected_delivery = room.get("expected_delivery_date")
            if expected_delivery:
                if isinstance(expected_delivery, str):
                    try:
                        date_obj = datetime.fromisoformat(expected_delivery.replace('Z', '+00:00'))
                        ws.cell(row=row_idx, column=9, value=date_obj.strftime("%Y-%m-%d"))
                    except:
                        ws.cell(row=row_idx, column=9, value=expected_delivery)
                else:
                    ws.cell(row=row_idx, column=9, value=str(expected_delivery))
            else:
                ws.cell(row=row_idx, column=9, value="")
            
            ws.cell(row=row_idx, column=10, value=room.get("letter_status", "") or "无")
        
        # 自动调整列宽
        print("6. 调整列宽...")
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 20)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # 保存到内存
        print("7. 保存到内存...")
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        # 生成文件名
        current_time = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"ZWY项目汇总_{current_time}.xlsx"
        
        print(f"成功生成Excel文件: {filename}")
        print(f"文件大小: {len(output.getvalue())} 字节")
        
    except Exception as e:
        print(f"导出过程中发生错误: {str(e)}")
        traceback.print_exc()
    finally:
        db.close()

if __name__ == "__main__":
    test_export_direct()